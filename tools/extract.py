import openpyxl, json, sys
wb=openpyxl.load_workbook(sys.argv[1],data_only=True); ws=wb['for copy']
v=lambda a: ws[a].value
d={}
# psychro tables
temps=[ws.cell(row=r,column=76).value for r in range(91,116)]  # BX=76
rhs=[ws.cell(row=90,column=c).value for c in range(77,89)]     # BY..CJ
d['psy']={'T':temps,'RH':rhs,
  'h':[[ws.cell(row=r,column=c).value for c in range(77,89)] for r in range(91,116)],
  'v':[[ws.cell(row=r,column=c).value for c in range(91,103)] for r in range(91,116)]}  # CM..CX = 91..102
assert ws.cell(row=91,column=90).value==-60 and ws.cell(row=90,column=91).value==0
# foods AI88:AO315
foods=[]
for r in range(88,316):
    n=ws.cell(row=r,column=35).value
    if n is None: continue
    foods.append([n]+[ws.cell(row=r,column=c).value for c in range(36,42)])
d['foods']=foods  # name, category, tf, c_above, c_below, moisture, latent
cats={}
for c in range(54,67):  # BB..BN
    name=ws.cell(row=88,column=c).value
    items=[ws.cell(row=r,column=c).value for r in range(89,149) if ws.cell(row=r,column=c).value not in (None,0)]
    cats[name]=items
d['categories']=cats
d['k']={ws.cell(row=r,column=17).value: ws.cell(row=r,column=18).value for r in range(100,105)}
d['thkPS']=[[v(f'FE{r}'),v(f'FF{r}')] for r in range(88,95)]
d['thkPU']=[[v(f'FH{r}'),v(f'FI{r}')] for r in range(88,93)]
d['roomTypes']=[v(f'FB{r}') for r in range(88,96)]
d['fanFactor']={v(f'FB{r}'):v(f'FC{r}') for r in range(88,96)}
d['lighting']={v(f'P{r}'):v(f'R{r}') for r in range(59,67)}
d['protection']=[[v(f'P{r}'),v(f'S{r}')] for r in range(19,27)]
abf=[]
for r in range(89,475):
    x=ws.cell(row=r,column=122).value; y=ws.cell(row=r,column=123).value  # DR=122, DS=123
    if isinstance(x,(int,float)) and isinstance(y,(int,float)) and not isinstance(x,bool): abf.append([x,y])
d['abf']=abf
json.dump(d,open(sys.argv[2],'w'),ensure_ascii=False)
print(len(foods),'foods', {k:len(x) for k,x in cats.items()}, 'abf pts',len(abf), d['k'], d['fanFactor'], d['lighting'])
