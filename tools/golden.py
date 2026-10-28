# Generate golden cases, fill V1.25 copies, recalc in LibreOffice, store expected values.
import openpyxl, json, random, subprocess, sys, os, glob
master, outdir, n = sys.argv[1], sys.argv[2], int(sys.argv[3])
data = json.load(open(os.path.join(os.path.dirname(__file__), '../src/data/v125.json')))
R = random.Random(2545)
types = data['roomTypes']; cats = data['categories']
mats = ['EPS','XPS','PU','PIR']
cols = ['D','G','J','M']; pcols = ['D','F','H','J','L']
def roomT(t):
    return {'Ante Room':R.choice([5,8,10,12]),'Chilled Room':R.choice([-2,0,2,4]),'Cold Room':R.choice([-25,-20,-18,0,2,4]),
            'Freezer - IQF':R.choice([-40,-35,-30]),'Freezer - Trolley':R.choice([-40,-35,-30]),'Loading Dock':R.choice([5,8,10,12]),
            'Processing Area':R.choice([10,12,15,18]),'Air Conditioning':R.choice([22,24,25,26])}[t]
cases=[]
for i in range(n):
    t = types[i % len(types)]
    T = roomT(t); mat = R.choice(mats)
    c = {'id':i,'roomType':t,'T':T,'RH':R.choice([50,60,72.5,85,90,95]),'H':R.choice([3,4,4.5,6,8]),'W':R.choice([4,6,8.5,12]),'L':R.choice([5,10,15,22]),
         'ceiling':{'T':R.choice([35,40,45]),'type':mat},'floor':{'T':R.choice([25,30]),'convection':R.choice(['Yes','No'])},
         'walls':[],'products':[],'workers':R.choice([0,1,2,4]),'safetyFactor':R.choice([10,15,20]),
         'makeUp':{'T':35,'RH':60,'CMH':R.choice([0,0,500,1200])},'equipment':[]}
    if R.random()<0.3: c['floor']['type']=R.choice(mats+['Concrete'])
    for w in range(4):
        wall={'T':R.choice([T+5, 20, 30, 35, 38, T-5]),'doors':[]}
        if R.random()<0.25: wall['type']=R.choice(mats+['Concrete'])
        for d in range(2):
            if R.random()<0.35:
                wall['doors'].append({'RH':R.choice([50,60,75,85]),'count':R.choice([1,2]),'E':R.choice([0,0.7,0.8,0.85,0.95]),
                                      'openMin':R.choice([0,5,10,15,30]),'W':R.choice([1.2,2.4,2.5,3.0]),'H':R.choice([2.1,2.4,3.0])})
            else: wall['doors'].append({})
        c['walls'].append(wall)
    for p in range(R.choice([0,1,2,3,5])):
        cat = R.choice(list(cats)); name = R.choice(cats[cat])
        tin = R.choice([T+20, T+10, 5, 0, -1, -2, -10, 10, 25]); tout = R.choice([T, T+2, -18, -1, 2])
        c['products'].append({'category':cat,'name':name,'tIn':tin,'tOut':tout,'ton':R.choice([0.5,2,5,12.5]),'hours':R.choice([4,8,12,24]),'respiration':R.choice([0,0,20,60])})
    if R.random()<0.5:
        c['forklift']={'kW':R.choice([5,7]),'hours':R.choice([4,8]),'kgPerPallet':R.choice([500,600,1000]),'tonPerDay':R.choice([5,20,60])}
    for e in range(R.choice([0,2,5,10])):
        c['equipment'].append({'kW':R.choice([0.5,1.5,3.7,7.5]),'factor':R.choice([0.3,1,0.5])})
    if R.random()<0.2: c['fanMotorKW']=R.choice([1.1,2.2,4.4])
    cases.append(c)

def fill(c, path):
    wb=openpyxl.load_workbook(master); ws=wb['for copy']
    S=lambda a,v: ws.__setitem__(a,v)
    S('J3',f"CASE {c['id']}"); S('J4',c['roomType']); S('J9',c['T']); S('J10',c['RH']); S('J12',c['H']); S('J13',c['W']); S('J14',c['L'])
    S('J17',c['makeUp']['T']); S('J18',c['makeUp']['RH']); S('J19',c['makeUp']['CMH'])
    S('M9',c['ceiling']['T']); S('M11',c['ceiling']['type']); S('M15',c['floor']['T']); S('M19',c['floor']['convection'])
    if 'type' in c['floor']: S('M17',c['floor']['type'])
    for wi,w in enumerate(c['walls']):
        col=cols[wi]; S(f'{col}22',w['T'])
        if 'type' in w: S(f'{col}23',w['type'])
        for di,d in enumerate(w['doors']):
            if not d: continue
            base = 29 if di==0 else 38
            S(f'{col}{base}',d['RH']); S(f'{col}{base+1}',d['count']); S(f'{col}{base+2}',d['E']); S(f'{col}{base+3}',d['openMin']); S(f'{col}{base+4}',d['W']); S(f'{col}{base+5}',d['H'])
    for pi,p in enumerate(c['products']):
        col=pcols[pi]; S(f'{col}48',p['category']); S(f'{col}49',p['name']); S(f'{col}50',p['tIn']); S(f'{col}51',p['tOut']); S(f'{col}52',p['ton']); S(f'{col}53',p['hours']); S(f'{col}58',p['respiration'])
    S('D63',c['workers']); S('L76',c['safetyFactor'])
    if 'fanMotorKW' in c: S('D66',c['fanMotorKW'])
    if 'forklift' in c:
        f=c['forklift']; S('G62',f['kW']); S('G63',f['hours']); S('G64',f['kgPerPallet']); S('G65',f['tonPerDay'])
    for ei,e in enumerate(c['equipment']):
        col,fc = ('I','J') if ei<5 else ('L','M'); r=62+ei%5
        S(f'{col}{r}',e['kW']); S(f'{fc}{r}',e['factor'])
    wb.save(path)

os.makedirs(outdir+'/in',exist_ok=True); os.makedirs(outdir+'/out',exist_ok=True)
for c in cases: fill(c, f"{outdir}/in/case{c['id']:03d}.xlsx")
subprocess.run(['soffice','--headless','--convert-to','xlsx:Calc MS Excel 2007 XML','--outdir',outdir+'/out']+sorted(glob.glob(outdir+'/in/*.xlsx')),capture_output=True)
cells=['C69','C70','C71','C72','C73','C74','C75','F69','F70','F71','F72','F73','F74','F75','F76','F77','F78','I69','I70','I71','I72','I73','I74','L69','L70','L71','L72','L73','L75','L77','L78','L79','G66']
exp=[]
for c in cases:
    ws=openpyxl.load_workbook(f"{outdir}/out/case{c['id']:03d}.xlsx",data_only=True)['for copy']
    exp.append({a:ws[a].value for a in cells})
json.dump({'cases':cases,'expected':exp},open(os.path.join(os.path.dirname(__file__),'../test/golden.json'),'w'),ensure_ascii=False,indent=0)
print('cases',len(cases))
